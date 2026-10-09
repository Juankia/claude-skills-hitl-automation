#!/usr/bin/env python3
"""Render a first-draft engagement letter from the shared client state.

Usage:
    python draft_document.py --client-id <id> [--format docx|md] [--out PATH]
                             [--template PATH] [--firm PATH] [--date YYYY-MM-DD]

The template is Markdown with a deliberately tiny syntax so the output is
predictable and reviewable:

    {{ client.legal_name }}               variable (dotted path)
    {{ engagement.retainer_usd | money }} filters: money, title, join, longdate, upper
    {% if engagement.fee_type == "hourly" %} ... {% endif %}
    {% if consent.electronic_communications %} / {% if not x %}

Unknown values never crash the render: they become ``[[MISSING: path]]``
markers (highlighted in the .docx) and are listed in the JSON result so a human
can fill them. Jurisdiction-specific clauses come from
``references/jurisdictions.json``; unknown states fall back to ``default`` and
produce ``[[REVIEW: ...]]`` markers.

The .docx output is byte-for-byte deterministic for the same inputs (fixed
document properties and zip timestamps), which makes drafts diffable and
lets tests assert on hashes.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
for _candidate in (HERE, HERE.parents[2] / "shared"):
    if (_candidate / "client_state.py").exists():
        sys.path.insert(0, str(_candidate))
        break

from client_state import ClientStateStore, StateError  # noqa: E402

DEFAULT_TEMPLATE = SKILL_DIR / "templates" / "engagement_letter.md"
DEFAULT_FIRM = SKILL_DIR / "templates" / "firm_profile.example.json"
JURISDICTIONS = SKILL_DIR / "references" / "jurisdictions.json"

VAR_RE = re.compile(r"\{\{\s*([A-Za-z_][\w.]*)\s*(?:\|\s*(\w+)\s*)?\}\}")
IF_RE = re.compile(r"^\{%\s*if\s+(not\s+)?([\w.]+)(?:\s*==\s*\"([^\"]*)\")?\s*%\}$")
ENDIF_RE = re.compile(r"^\{%\s*endif\s*%\}$")
MARKER_RE = re.compile(r"\[\[(MISSING|REVIEW):[^\]]*\]\]")
_MISSING = object()


class TemplateError(ValueError):
    pass


# --- context ---------------------------------------------------------------

def load_jurisdiction(code: str) -> tuple[dict[str, Any], bool]:
    table = json.loads(JURISDICTIONS.read_text(encoding="utf-8"))
    if code in table:
        return table[code], True
    return table["default"], False


def build_context(record: dict[str, Any], firm: dict[str, Any], today: date) -> tuple[dict[str, Any], list[str]]:
    notes: list[str] = []
    jurisdiction, known = load_jurisdiction(record["matter"]["jurisdiction"])
    if not known:
        notes.append(
            f"No clause set for jurisdiction {record['matter']['jurisdiction']}; default language used."
        )
    client = dict(record["client"])
    client["salutation"] = client.get("preferred_name") or client["legal_name"]
    if record.get("intake", {}).get("status") == "needs_review":
        notes.extend(f"Intake review: {r}" for r in record["intake"].get("review_reasons", []))
    context = {
        **record,
        "client": client,
        "firm": firm,
        "jurisdiction": jurisdiction,
        "today": today.isoformat(),
    }
    return context, notes


# --- tiny template engine --------------------------------------------------

def lookup(context: dict[str, Any], path: str) -> Any:
    value: Any = context
    for part in path.split("."):
        if isinstance(value, dict) and part in value:
            value = value[part]
        else:
            return _MISSING
    return value


def _money(v: Any) -> str:
    return f"${float(v):,.2f}"


def _join(v: Any) -> str:
    items = [str(i) for i in v]
    if len(items) <= 2:
        return " and ".join(items)
    return ", ".join(items[:-1]) + ", and " + items[-1]


def _longdate(v: Any) -> str:
    d = date.fromisoformat(str(v))
    return f"{d:%B} {d.day}, {d.year}"


FILTERS = {
    "money": _money,
    "title": lambda v: str(v).replace("_", " ").title(),
    "join": _join,
    "longdate": _longdate,
    "upper": lambda v: str(v).upper(),
}


def _truthy(context: dict[str, Any], negate: bool, path: str, literal: str | None) -> bool:
    value = lookup(context, path)
    if literal is not None:
        result = value is not _MISSING and str(value) == literal
    else:
        result = value is not _MISSING and bool(value)
    return not result if negate else result


def render(template: str, context: dict[str, Any]) -> tuple[str, list[str]]:
    """Return (rendered_markdown, sorted list of missing variable paths)."""
    missing: set[str] = set()
    out: list[str] = []
    stack: list[bool] = []  # visibility of each open if-block

    for lineno, line in enumerate(template.splitlines(), start=1):
        stripped = line.strip()
        if m := IF_RE.match(stripped):
            stack.append(_truthy(context, bool(m.group(1)), m.group(2), m.group(3)))
            continue
        if ENDIF_RE.match(stripped):
            if not stack:
                raise TemplateError(f"line {lineno}: endif without if")
            stack.pop()
            continue
        if stripped.startswith("{%"):
            raise TemplateError(f"line {lineno}: unsupported tag {stripped!r}")
        if not all(stack):
            continue

        def substitute(match: re.Match[str]) -> str:
            path, flt = match.group(1), match.group(2)
            value = lookup(context, path)
            if value is _MISSING or value is None or value == "":
                missing.add(path)
                return f"[[MISSING: {path}]]"
            if flt:
                if flt not in FILTERS:
                    raise TemplateError(f"line {lineno}: unknown filter {flt!r}")
                return FILTERS[flt](value)
            return str(value)

        out.append(VAR_RE.sub(substitute, line))

    if stack:
        raise TemplateError("unterminated {% if %} block")
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"
    return text, sorted(missing)


# --- Markdown -> .docx -----------------------------------------------------

FIXED_TIMESTAMP = datetime(2026, 1, 1, tzinfo=timezone.utc)


def _add_runs(paragraph, text: str) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    # Split on **bold** and [[MARKERS]], keeping the delimiters.
    for token in re.split(r"(\*\*[^*]+\*\*|\[\[[^\]]+\]\])", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif MARKER_RE.fullmatch(token):
            run = paragraph.add_run(token)
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            paragraph.add_run(token)


def markdown_to_docx(markdown: str, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    props = doc.core_properties
    props.title = title
    props.author = "document-drafting skill"
    props.last_modified_by = "document-drafting skill"
    props.comments = "FIRST DRAFT - requires professional review before sending."
    props.created = props.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    props.revision = 1

    block: list[str] = []

    def flush() -> None:
        if not block:
            return
        para = doc.add_paragraph()
        for i, line in enumerate(block):
            if i:
                para.add_run().add_break()
            _add_runs(para, line)
        block.clear()

    for line in markdown.splitlines():
        if not line.strip():
            flush()
        elif line.startswith("# "):
            flush()
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            flush()
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            flush()
            _add_runs(doc.add_paragraph(style="List Bullet"), line[2:].strip())
        else:
            block.append(line.rstrip())
    flush()

    raw = io.BytesIO()
    doc.save(raw)
    return _normalize_zip(raw.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """Rewrite the OOXML zip with fixed timestamps so output is reproducible."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


# --- orchestration ---------------------------------------------------------

def draft(
    client_id: str,
    store: ClientStateStore,
    fmt: str = "docx",
    out: Path | None = None,
    template_path: Path = DEFAULT_TEMPLATE,
    firm_path: Path = DEFAULT_FIRM,
    today: date | None = None,
) -> dict[str, Any]:
    record = store.get_client(client_id)
    firm = json.loads(firm_path.read_text(encoding="utf-8"))
    context, notes = build_context(record, firm, today or date.today())
    markdown, missing = render(template_path.read_text(encoding="utf-8"), context)
    review_markers = sorted(set(re.findall(r"\[\[REVIEW:\s*([^\]]+)\]\]", markdown)))

    out = out or Path("out") / client_id / f"engagement-letter-{client_id}.{fmt}"
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "md":
        payload = markdown.encode("utf-8")
    else:
        payload = markdown_to_docx(markdown, title=f"Engagement Letter - {record['client']['legal_name']}")
    out.write_bytes(payload)
    digest = hashlib.sha256(payload).hexdigest()

    result = {
        "ok": True,
        "client_id": client_id,
        "output": str(out),
        "format": fmt,
        "sha256": digest,
        "jurisdiction": record["matter"]["jurisdiction"],
        "missing": missing,
        "review_markers": review_markers,
        "notes": notes,
        "status": "draft_needs_attention" if (missing or review_markers or notes) else "draft_ready_for_review",
    }
    store.add_artifact(
        client_id,
        "documents",
        {"type": "engagement_letter", "path": str(out), "sha256": digest, "status": result["status"]},
    )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Draft an engagement letter from the shared client state.")
    parser.add_argument("--client-id", required=True)
    parser.add_argument("--format", choices=["docx", "md"], default="docx")
    parser.add_argument("--out", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--firm", type=Path, default=DEFAULT_FIRM)
    parser.add_argument("--state-dir", type=Path)
    parser.add_argument("--date", type=date.fromisoformat, help="letter date (default: today)")
    args = parser.parse_args(argv)

    try:
        result = draft(
            args.client_id,
            ClientStateStore(args.state_dir),
            fmt=args.format,
            out=args.out,
            template_path=args.template,
            firm_path=args.firm,
            today=args.date,
        )
    except (StateError, TemplateError, OSError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, indent=2))
        return 1
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
